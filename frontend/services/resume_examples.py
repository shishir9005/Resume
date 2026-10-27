"""
Example resumes for the Resources page.

The examples are Markdown files in backend/knowledge/ (example-resume-*.md), so the
Resume Coach can retrieve them too. This module loads them for display and turns
each one into a single-column, ATS-friendly .docx template.
"""
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

KNOWLEDGE_DIR = Path(__file__).resolve().parents[2] / "backend" / "knowledge"

# (file name, tab label, score this app gave the example when it was written, Oct 2026).
# LLM parsing varies a little between runs, so treat the scores as approximate.
EXAMPLES = [
    ("example-resume-backend", "💻 Backend Engineering",  89),
    ("example-resume-ai",      "🤖 AI / LLM Engineering", 91),
    ("example-resume-data",    "📊 Data Analyst",         88),
]


@dataclass
class ResumeExample:
    slug: str
    label: str
    score: int
    title: str      # the '# ' line, e.g. "Example Resume: Backend Engineering Intern"
    body: str       # everything after the title, as Markdown


def load_examples() -> List[ResumeExample]:
    examples = []
    for slug, label, score in EXAMPLES:
        path = KNOWLEDGE_DIR / f"{slug}.md"
        if not path.exists():
            continue
        title_line, _, body = path.read_text(encoding="utf-8").partition("\n")
        examples.append(ResumeExample(slug, label, score, title_line.lstrip("# ").strip(), body.strip()))
    return examples


def _add_bottom_border(paragraph) -> None:
    """Thin line under a section heading. Pure paragraph formatting, so ATS parsers ignore it."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")):
        bottom.set(qn(f"w:{attr}"), value)
    borders.append(bottom)
    p_pr.append(borders)


def _add_rich_text(paragraph, text: str) -> None:
    """Write text into a paragraph, turning **bold** spans into bold runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = (i % 2 == 1)


def to_docx(example: ResumeExample) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    lines = [line.rstrip() for line in example.body.splitlines()]
    name_done = False

    for line in lines:
        if not line:
            continue

        # first bold-only line is the candidate's name
        if not name_done and re.fullmatch(r"\*\*(.+)\*\*", line):
            run = doc.add_paragraph().add_run(line.strip("*"))
            run.bold = True
            run.font.size = Pt(18)
            name_done = True
        elif line.startswith("## "):
            heading = doc.add_paragraph()
            heading.paragraph_format.space_before = Pt(8)
            run = heading.add_run(line[3:].upper())
            run.bold = True
            run.font.size = Pt(11.5)
            _add_bottom_border(heading)
        elif line.startswith("- "):
            # a plain "•" character keeps bullets visible to every ATS parser
            bullet = doc.add_paragraph()
            bullet.paragraph_format.left_indent = Inches(0.2)
            bullet.paragraph_format.first_line_indent = Inches(-0.15)
            _add_rich_text(bullet, "• " + line[2:])
        else:
            _add_rich_text(doc.add_paragraph(), line)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
