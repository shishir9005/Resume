import io

from docx import Document

from frontend.services.resume_examples import EXAMPLES, load_examples, to_docx


def test_every_example_loads():
    examples = load_examples()
    assert [e.slug for e in examples] == [slug for slug, _, _ in EXAMPLES]
    assert all(e.body.startswith('**') for e in examples)   # body opens with the candidate's name


def test_docx_template_keeps_all_content():
    example = load_examples()[0]
    doc = Document(io.BytesIO(to_docx(example)))
    text = '\n'.join(p.text for p in doc.paragraphs)

    assert doc.paragraphs[0].text == 'Lin Wei'
    for heading in ('SUMMARY', 'EDUCATION', 'PROJECTS', 'EXPERIENCE', 'SKILLS'):
        assert heading in text
    assert '**' not in text                    # markdown bold became real bold runs
    assert text.count('• ') >= 10              # bullets survive as plain characters
    assert not doc.tables                      # single column, no tables
